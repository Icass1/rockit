"""Keep radio UI vocabulary in the workbook and generated shared contract."""

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[4]


def test_radio_vocabulary_is_in_workbook_and_generated_types() -> None:
    """Every static vocabulary reference in the radio experience must exist."""
    workbook = openpyxl.load_workbook(ROOT / "Vocabulary.xlsx", data_only=True)
    try:
        sheet = workbook.active
        assert sheet is not None
        keys = {
            row[0]
            for row in sheet.iter_rows(min_row=3, values_only=True)
            if isinstance(row[0], str)
        }
    finally:
        workbook.close()

    generated = (ROOT / "packages/shared/src/models/types/vocabulary.ts").read_text()
    generated_keys = set(re.findall(r"^    ([A-Z][A-Z0-9_]+): string;", generated, re.M))
    assert generated_keys == keys

    web = ROOT / "apps/web"
    sources = [
        *sorted((web / "components/Radio").glob("*.tsx")),
        web / "components/Footer/FooterCenter.tsx",
        web / "components/Footer/FooterLeft.tsx",
        web / "lib/managers/stationManager.ts",
    ]
    references: set[str] = set()
    for source in sources:
        references.update(
            re.findall(
                r"\b(?:v|vocabulary|\$vocabulary)\.([A-Z][A-Z0-9_]+)",
                source.read_text(),
            )
        )
    assert references <= keys, sorted(references - keys)


def test_new_radio_translations_are_importable_without_excel_recalculation() -> None:
    """The importer reads cached values; formulas alone would silently lose text."""
    required = {
        "RADIO_DESCRIPTION",
        "RADIO_EXPLORE",
        "RADIO_SAVED",
        "RADIO_COUNTRY",
        "RADIO_GRID",
        "RADIO_LIST",
        "RADIO_CLEAR",
        "RADIO_LOADING",
        "RADIO_LOAD_ERROR",
        "RADIO_UNAVAILABLE",
        "RADIO_MAP_HINT",
        "RADIO_EMPTY",
        "PLAYER_ENABLE_SHUFFLE",
        "PLAYER_DISABLE_SHUFFLE",
    }
    workbook = openpyxl.load_workbook(ROOT / "Vocabulary.xlsx", data_only=True)
    try:
        sheet = workbook.active
        assert sheet is not None
        languages = [cell.value for cell in sheet[2]][1:]
        assert languages == ["en", "es", "fr"]
        rows = {
            row[0]: row[1:]
            for row in sheet.iter_rows(min_row=3, values_only=True)
            if row[0] in required
        }
        assert rows.keys() == required
        for key, translations in rows.items():
            assert all(isinstance(text, str) and text.strip() for text in translations), key
    finally:
        workbook.close()
